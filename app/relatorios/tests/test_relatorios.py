"""
Testes do módulo de relatórios (não precisam do Firebird nem do backend).

O backend (app.config, app.rh.*, app.afastamentos.*, app.audit...) é substituído por
módulos falsos com dados de exemplo. Rodar a partir da pasta que contém `relatorios`:
    uv run pytest relatorios/tests -q
"""
import io
import sys
import types
from contextlib import contextmanager
from datetime import date, datetime, timezone
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

D = date
AGORA = datetime(2026, 9, 29, 12, 0, tzinfo=timezone.utc)

# ============================================================================ backend falso
USUARIO = {"id": 1, "email": "rh@caceres.mt.gov.br", "permissions": []}
AUDITORIA: list = []
PERFIS = {
    "admin": ["rh:diagnostico:painel", "rh:diagnostico:listar", "rh:servidores:listar", "rh:servidores:detalhar"],
    "rh": ["rh:diagnostico:painel", "rh:diagnostico:listar", "rh:servidores:listar", "rh:servidores:detalhar"],
    "consulta": ["rh:diagnostico:painel", "rh:servidores:listar"],
    "listar_sem_cpf": ["rh:diagnostico:painel", "rh:diagnostico:listar", "rh:servidores:listar"],
}


def _vinc(reg, mat, nome, cpf, sit, sit_nome, div, div_nome, cargo, cargo_nome, contrato=1):
    return {"REGISTRO": reg, "MATRICULA": mat, "CONTRATO": contrato, "NOME": nome, "CPF": cpf,
            "SITUACAO": sit, "SITUACAO_NOME": sit_nome, "VINCULO": "02", "VINCULO_NOME": "EFETIVO",
            "CARGOATUAL": cargo, "CARGOATUAL_NOME": cargo_nome, "DIVISAO": div, "DIVISAO_NOME": div_nome,
            "SUBDIVISAO": None, "SUBDIVISAO_NOME": None, "LOCAL_TRABALHO": "010", "LOCAL_TRABALHO_NOME": "ESCOLA X",
            "DTADMISSAO": datetime(2010, 3, 1), "DTDEMISSAO": None}


EDU, SAU = "000003", "000010"
VINCULOS_SIP = [
    _vinc("028310", 4618, "ODENISE JARA GOMES LENTE", "53175611168", 1, "Normal", EDU,
          "SECRETARIA MUNICIPAL DE EDUCAÇÃO", "101", "PROF LIC PEDAG C/DOC (30HS/AULA)"),
    _vinc("028312", 4618, "ODENISE JARA GOMES LENTE", "53175611168", 1, "Normal", EDU,
          "SECRETARIA MUNICIPAL DE EDUCAÇÃO", "102", "PROF LIC PEDAG C/ SUPERV(30HS/AULA)", contrato=2),
    _vinc("029982", 24935, "NATÁLIA LOPES RODRIGUES SIQUEIRA", "05271934101", 1, "Normal", "000020",
          "SECRETARIA MUNICIPAL DE ASSISTÊNCIA SOCIAL E CIDADANIA", "103", "PSICOLOGO N/S (40HS)"),
    _vinc("022818", 11669, "JOAO GONÇALO DO NASCIMENTO", "02511951126", 1, "Normal", EDU,
          "SECRETARIA MUNICIPAL DE EDUCAÇÃO", "101", "PROF LIC PEDAG C/DOC (30HS/AULA)"),
    _vinc("016074", 11448, "JOELI CASSIA DA SILVA MAGALHÃES", "03534387155", 1, "Normal", SAU,
          "SECRETARIA MUNICIPAL DE SAÚDE", "104", "TECNICO EM ENFERMAGEM(S/G.T)"),
    _vinc("000099", 99, "ALVARO DESLIGADO", "11111111111", 6, "Desligado", SAU,
          "SECRETARIA MUNICIPAL DE SAÚDE", "104", "TECNICO EM ENFERMAGEM(S/G.T)"),
]


class Regra:
    def __init__(self, codigo, titulo, tipo, alvo):
        self.codigo, self.titulo, self.tipo, self.alvo = codigo, titulo, tipo, alvo
        self.descricao, self.orientacao = f"Descrição {codigo}", f"Orientação {codigo}"


REGRAS = [Regra("S01", "Servidor sem celular", "alerta", "servidor"),
          Regra("S10", "CPF inválido", "erro", "servidor"),
          Regra("D03", "IRRF acima da idade", "erro", "dependente"),
          Regra("I01", "Indicador qualquer", "indicador", "servidor")]


def _ocorrencia(reg, **extra):
    v = next(x for x in VINCULOS_SIP if x["REGISTRO"] == reg)
    base = {"REGISTRO": reg, "MATRICULA": v["MATRICULA"], "CONTRATO": v["CONTRATO"], "SERVIDOR": v["NOME"],
            "CPF_SERVIDOR": v["CPF"], "ADMISSAO": "01/03/2010", "SITUACAO": v["SITUACAO_NOME"],
            "VINCULO": "EFETIVO", "CARGO": v["CARGOATUAL_NOME"], "DIVISAO": v["DIVISAO_NOME"],
            "LOCAL_TRABALHO": "ESCOLA X", "CELULAR": "", "EMAIL": "", "DETALHE": "detalhe"}
    base.update(extra)
    return base


ATIVOS = [v for v in VINCULOS_SIP if v["SITUACAO"] == 1]
RESULTADO_DIAG = {
    "gerado_em": AGORA,
    "servidores": {v["REGISTRO"]: {**v} for v in ATIVOS},
    "classe": {"028310": "erro", "028312": "ok", "029982": "alerta", "022818": "erro", "016074": "ok"},
    "por_regra": {
        "S01": [_ocorrencia("029982")],
        "S10": [_ocorrencia("028310"), _ocorrencia("022818")],
        "D03": [_ocorrencia("022818", ITEM=1, DEPENDENTE="FILHO DO JOÃO", PARENTESCO="Filho",
                            NASCIMENTO="10/05/2001", IDADE=25, CPF_DEPENDENTE="", DEPENDENTE_IRRF="Sim")],
        "I01": [_ocorrencia("016074")],
    },
    "deps_por_registro": {"022818": 2},
}


def _af_item(reg, reduzido, tipo, dtini, dtfim, status="sem_status", causa="91", **extra):
    v = next(x for x in VINCULOS_SIP if x["REGISTRO"] == reg)
    item = {"registro": reg, "matricula": v["MATRICULA"], "nome": v["NOME"], "cpf": v["CPF"],
            "situacao_vinculo": 1, "situacao_vinculo_nome": "Normal",
            "divisao_codigo": v["DIVISAO"], "divisao": v["DIVISAO_NOME"],
            "cargo_codigo": v["CARGOATUAL"], "cargo": v["CARGOATUAL_NOME"],
            "item": 1, "reduzido": reduzido, "tipo": tipo, "causa": causa, "dtini": dtini, "dtfim": dtfim,
            "retorno_previsto": dtfim, "duracao_dias": (dtfim - dtini).days, "sem_remuneracao": causa == "97",
            "status_esocial": status, "documento_tipo": "ATESTADO MEDICO", "documento_numero": None,
            "motivo_codigo": None}
    item.update(extra)
    return item


VIGENTES = [
    _af_item("028310", 12, "AFASTAMENTO LICENCA MEDICA-EFETIVOS", D(2026, 6, 1), D(2026, 11, 14)),
    _af_item("028312", 12, "AFASTAMENTO LICENCA MEDICA-EFETIVOS", D(2026, 6, 1), D(2026, 11, 14)),
    _af_item("029982", 168, "LICENCA MATERNIDADE", D(2026, 7, 1), D(2026, 11, 27), status="valido", causa="50"),
    _af_item("022818", 378, "LICENCA S/REMUN. (VACANCIA)", D(2025, 3, 10), D(2028, 3, 9), status="valido", causa="97"),
]


def _resp(itens, **extra):
    return {"data_referencia": D(2026, 9, 29), "total": len(itens), "itens": itens, **extra}


class ServicoFalso(types.ModuleType):
    def data_da_copia(self, con):
        return {"data_copia": D(2026, 9, 29)}

    def listar_vigentes(self, con, data_ref=None, **_):
        return _resp([dict(i) for i in VIGENTES])

    def listar_retornos(self, con, dias=30, data_ref=None):
        return _resp([{**VIGENTES[2], "dias_ate_retorno": 59, "emenda_outro_afastamento": False}])

    def afastamentos_longos(self, con, min_dias=60, somente_sem_remuneracao=False, data_ref=None):
        return _resp([{"registro": "022818", "matricula": 11669, "nome": "JOAO GONÇALO DO NASCIMENTO",
                       "cpf": "02511951126", "divisao": "SECRETARIA MUNICIPAL DE EDUCAÇÃO", "cargo": "PROF",
                       "afastado_desde": D(2025, 3, 10), "dias_continuos": 569, "sem_remuneracao": True,
                       "afastamentos_vigentes": [{"reduzido": 378, "tipo": "LICENCA S/REMUN."}]}])

    def acumulado_no_ano(self, con, ano=None, min_dias=1, data_ref=None):
        return _resp([{"registro": "028310", "matricula": 4618, "nome": "ODENISE JARA GOMES LENTE",
                       "cpf": "53175611168", "divisao": "SECRETARIA MUNICIPAL DE EDUCAÇÃO", "cargo": "PROF",
                       "dias_afastado": 121, "afastamentos": 1, "dias_por_tipo": {"12 LIC": 121}}], ano=2026)

    def recorrentes(self, con, ano=None, min_episodios=4, **_):
        return _resp([{"registro": "016074", "matricula": 11448, "nome": "JOELI CASSIA DA SILVA MAGALHÃES",
                       "cpf": "03534387155", "divisao": "SECRETARIA MUNICIPAL DE SAÚDE", "cargo": "TEC",
                       "episodios": 6, "lancamentos": 6, "dias_afastado": 9, "fator_bradford": 324,
                       "episodios_curtos": 5, "maior_episodio_dias": 2, "perfil": "frequente",
                       "tipos": {"12 LIC": 6}, "lista_episodios": [{"dias": 1}]}],
                     ano=2026, total_por_perfil={"frequente": 1}, legenda={})

    def esocial_pendencias(self, con, status, desde=None, reduzido=None):
        return {"desde": desde, "total": 2, "itens": [dict(VIGENTES[0]), dict(VIGENTES[1])]}

    def inconsistencias(self, con, regra, desde=None, data_ref=None):
        itens = []
        if regra == "sobrepostos":
            itens = [{"registro": "028310", "matricula": 4618, "nome": "ODENISE JARA GOMES LENTE",
                      "divisao": "SECRETARIA MUNICIPAL DE EDUCAÇÃO", "a_item": 1, "a_reduzido": 12, "a_tipo": "LIC",
                      "a_dtini": D(2026, 1, 1), "a_dtfim": D(2026, 1, 11), "b_item": 2, "b_reduzido": 228,
                      "b_tipo": "AUX", "b_dtini": D(2026, 1, 10), "b_dtfim": D(2026, 1, 20), "dias_sobrepostos": 1}]
        return _resp(itens, regra=regra, descricao="desc")

    def abate_tempo(self, con, min_dias=1, data_ref=None):
        return _resp([])

    def maternidade_sem_dependente(self, con, desde=None, data_ref=None):
        return _resp([{**VIGENTES[2], "licenca": "maternidade", "qtd_dependentes": 0}])

    def resumir_vigentes(self, con, data_ref=None, por="tipo"):
        return {"data_referencia": D(2026, 9, 29), "agrupado_por": por, "total_vinculos": 4, "total_afastamentos": 4,
                "itens": [{"codigo": 12, "nome": "LIC MEDICA", "vinculos": 2, "afastamentos": 2}]}

    def absenteismo(self, con, ano=None, mes=None, data_ref=None):
        linha = {"divisao_codigo": EDU, "divisao": "EDUCAÇÃO", "vinculos": 3, "vinculos_afastados": 2,
                 "dias_afastados": 58, "percentual_dias_afastados": 7.7}
        return {"data_referencia": D(2026, 9, 29), "periodo": {}, "total": dict(linha), "itens": [linha]}

    def serie_mensal(self, con, meses=12, data_ref=None, divisao=None, reduzido=None):
        return {"data_referencia": D(2026, 9, 29), "itens": [
            {"mes": "2026-08", "afastamentos_iniciados": 3, "vinculos_afastados_no_mes": 4,
             "vinculos_afastados_no_fim_do_mes": 4, "dias_afastados": 100, "mes_incompleto": False}]}

    def esocial_tipos_com_problema(self, con, desde=None):
        return {"total": 1, "itens": [{"reduzido": 12, "tipo": "LIC", "problema": "sem_motivo_configurado",
                                       "lancamentos_desde": 3577, "ultimo_uso": D(2026, 9, 25)}]}


def _instalar_backend():
    def modulo(nome, **attrs):
        m = types.ModuleType(nome)
        m.__dict__.update(attrs)
        sys.modules[nome] = m
        return m

    app = modulo("app")
    app.__path__ = []

    def registrar(db, request, acao, user_id=None, email=None, target=None, detail=None):
        AUDITORIA.append({"acao": acao, "alvo": target, "detalhe": detail})

    modulo("app.audit", registrar=registrar)

    def get_db():
        yield None

    modulo("app.database", get_db=get_db)
    modulo("app.config", settings=types.SimpleNamespace(
        diagnostico_cache_minutes=10, situacoes_ativas=["1", "3"], sip_empresa="001"))
    rh = modulo("app.rh")
    rh.__path__ = []

    def current_user():
        return USUARIO

    modulo("app.rh.deps", current_user=current_user)
    modulo("app.rh.firebird", query=lambda sql, params=(): [dict(v) for v in VINCULOS_SIP])
    modulo("app.rh.regras", REGRAS=REGRAS, POR_CODIGO={r.codigo: r for r in REGRAS})

    def agregar_por_divisao(dados):
        linha = {"codigo": EDU, "nome": "EDUCAÇÃO", "vinculos": 3, "com_erro": 2, "com_alerta": 0,
                 "sem_pendencia": 1, "dependentes": 2, "conformidade": 33.3}
        return {"gerado_em": AGORA, "divisoes": [linha], "total": {**linha, "codigo": "ALL", "nome": "Toda"}}

    modulo("app.rh.diagnostico", resultado=lambda atualizar=False: RESULTADO_DIAG,
           agregar_por_divisao=agregar_por_divisao)

    def montar(linhas, hoje, divisao=None, subdivisao=None, vinculo=None):
        faixa = {"faixa": "0-13", "idade_minima": 0, "idade_maxima": 13, "exige_comprovacao": False,
                 "dependentes": 5, "irrf": 2, "salario_familia": 1, "universitario": 0, "deficiente": 0,
                 "irrf_fora_do_limite": 0, "salario_familia_fora_do_limite": 0, "em_desconformidade": 0}
        return {"faixas": [faixa], "dependentes": 5, "em_desconformidade": 0}

    modulo("app.rh.faixas", dados=lambda atualizar=False: {"linhas": [], "gerado_em": AGORA}, montar=montar)
    af = modulo("app.afastamentos")
    af.__path__ = []
    sys.modules["app.afastamentos.servico"] = ServicoFalso("app.afastamentos.servico")

    @contextmanager
    def conexao_somente_leitura():
        yield object()

    modulo("app.afastamentos.db", conexao_somente_leitura=conexao_somente_leitura)


_instalar_backend()

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from relatorios import busca as B  # noqa: E402
from relatorios import fontes  # noqa: E402
from relatorios.catalogo import CATALOGO, POR_ID  # noqa: E402
from relatorios.router import router  # noqa: E402

app = FastAPI()
app.include_router(router)


@pytest.fixture
def cliente():
    fontes.limpar_cache()
    AUDITORIA.clear()
    return TestClient(app)


def como(perfil):
    USUARIO["permissions"] = PERFIS[perfil]


# ============================================================================ busca
def test_normalizar_tira_acento_e_pontuacao():
    assert B.normalizar("NATÁLIA") == "natalia"
    assert B.normalizar("JOAO GONÇALO") == "joao goncalo"
    assert B.normalizar("PROF LIC PEDAG C/DOC (30HS/AULA)") == "prof lic pedag c doc 30hs aula"


@pytest.mark.parametrize("termo,esperado", [
    ("natalia", {"029982"}),
    ("goncalo", {"022818"}),
    ("magalhaes", {"016074"}),
    ("Magalhães", {"016074"}),
    ("odenise educacao", {"028310", "028312"}),        # palavras em campos diferentes
    ("4618", {"028310", "028312"}),                    # matrícula: a mesma pessoa, duas linhas
    ("28310", {"028310"}),                             # REGISTRO sem os zeros
    ("531.756.111-68", {"028310", "028312"}),          # CPF com pontuação
    ("53175611168", {"028310", "028312"}),             # CPF só números
    ("531756", {"028310", "028312"}),                  # começo do CPF
    ("c/doc", {"028310", "022818"}),
])
def test_busca(termo, esperado):
    linhas = [fontes.linha_vinculo(v, {"1", "3"}) for v in VINCULOS_SIP]
    achados = B.filtrar(linhas, termo, ["nome", "cargo", "secretaria"], pode_ver_cpf=True)
    assert {l["registro"] for l in achados} == esperado


def test_sem_permissao_de_cpf_a_busca_por_cpf_nao_acha():
    linhas = [fontes.linha_vinculo(v, {"1", "3"}) for v in VINCULOS_SIP]
    assert B.filtrar(linhas, "531.756.111-68", ["nome"], pode_ver_cpf=False) == []
    assert {l["registro"] for l in B.filtrar(linhas, "4618", ["nome"], pode_ver_cpf=False)} == {"028310", "028312"}


def test_trecho_curto_de_cpf_nao_vira_busca_por_cpf():
    linhas = [fontes.linha_vinculo(v, {"1", "3"}) for v in VINCULOS_SIP]
    assert B.filtrar(linhas, "531", ["nome"], pode_ver_cpf=True) == []


def test_ordenar_sem_acento_e_vazios_no_fim():
    linhas = [{"n": "Ética"}, {"n": None}, {"n": "Amanda"}, {"n": "Álvaro"}, {"n": "Bruno"}]
    assert [l["n"] for l in B.ordenar(linhas, "n")] == ["Álvaro", "Amanda", "Bruno", "Ética", None]
    assert [l["n"] for l in B.ordenar(linhas, "n", True)] == ["Ética", "Bruno", "Amanda", "Álvaro", None]


def test_agrupar_por_pessoa():
    linhas = [fontes.linha_vinculo(v, {"1", "3"}) for v in VINCULOS_SIP[:2]]
    [p] = B.agrupar_por_pessoa(linhas, ["registro", "nome", "cargo", "secretaria"])
    assert p["vinculos"] == 2
    assert p["registro"] == "028310 | 028312"
    assert p["secretaria"] == "SECRETARIA MUNICIPAL DE EDUCAÇÃO"     # igual nas duas: fica um só
    assert " | " in p["cargo"]


# ============================================================================ catálogo
def test_catalogo_coerente():
    assert len(POR_ID) == len(CATALOGO)
    for rel in CATALOGO:
        ids = {c.id for c in rel.colunas}
        assert len(ids) == len(rel.colunas), rel.id
        if rel.ordem_padrao:
            assert rel.ordem_padrao[0] in ids, rel.id
        assert rel.natureza in ("nominal", "agregado")
        if rel.nominal:
            assert "registro" in ids and "nome" in ids, rel.id
        for f in rel.filtros:
            assert f.tipo in ("selecao", "multipla", "numero", "data", "booleano"), (rel.id, f.id)
    for regra, padrao in fontes.COLUNAS_INCONSISTENCIA.items():
        assert set(padrao) <= {c.id for c in POR_ID["afastamentos-inconsistencias"].colunas}, regra


def test_catalogo_por_perfil(cliente):
    como("consulta")
    r = cliente.get("/rh/relatorios").json()
    ids = {x["id"] for x in r["relatorios"]}
    assert "vinculos" in ids and "diagnostico-secretarias" in ids
    assert "afastamentos-vigentes" not in ids and "diagnostico-regra" not in ids
    assert r["pode_ver_dados_sensiveis"] is False
    vinc = next(x for x in r["relatorios"] if x["id"] == "vinculos")
    assert "cpf" not in {c["id"] for c in vinc["colunas"]}

    como("rh")
    r = cliente.get("/rh/relatorios").json()
    assert len(r["relatorios"]) == len(CATALOGO)
    regra = next(x for x in r["relatorios"] if x["id"] == "diagnostico-regra")
    assert regra["filtros"][0]["opcoes"][0] == {"valor": "S01", "rotulo": "S01 — Servidor sem celular"}


def test_permissao_e_relatorio_inexistente(cliente):
    como("consulta")
    assert cliente.get("/rh/relatorios/afastamentos-vigentes").status_code == 403
    assert cliente.get("/rh/relatorios/nao-existe").status_code == 422      # id fora da lista


# ============================================================================ vínculos
def test_vinculos_padrao_so_ativos_e_opcoes(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/vinculos").json()
    assert r["total"] == 5 and r["total_carregado"] == 6
    assert [i["nome"] for i in r["itens"]][0] == "JOAO GONÇALO DO NASCIMENTO"     # ordem por nome
    sec = {o["valor"]: o["total"] for o in r["opcoes"]["secretaria"]}
    assert sec[EDU] == 3 and sec[SAU] == 1          # o desligado não conta: "Somente ativos" vale para as opções
    todos = cliente.get("/rh/relatorios/vinculos?ativo=false").json()
    assert todos["total"] == 6


def test_vinculos_filtro_secretaria_aceita_codigo_sem_zeros(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/vinculos?secretaria=3").json()
    assert r["total"] == 3
    r = cliente.get("/rh/relatorios/vinculos?secretaria=3,10").json()
    assert r["total"] == 4                                   # o desligado da Saúde fica de fora
    r = cliente.get("/rh/relatorios/vinculos?secretaria=3&secretaria=10&ativo=false").json()
    assert r["total"] == 5


def test_vinculos_busca_cpf_e_agrupar(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/vinculos", params={"q": "531.756.111-68"}).json()
    assert r["total"] == 2 and r["itens"][0]["cpf"] == "53175611168"
    g = cliente.get("/rh/relatorios/vinculos", params={"q": "odenise", "agrupar": "true"}).json()
    assert g["total"] == 1 and g["itens"][0]["vinculos"] == 2
    assert g["colunas"][1]["id"] == "vinculos"


def test_perfil_sem_cpf_nao_recebe_cpf_nem_acha_por_cpf(cliente):
    como("listar_sem_cpf")
    r = cliente.get("/rh/relatorios/vinculos", params={"q": "53175611168"}).json()
    assert r["total"] == 0
    r = cliente.get("/rh/relatorios/vinculos", params={"colunas": "registro,nome,cpf"}).json()
    assert all("cpf" not in i for i in r["itens"])
    assert [c["id"] for c in r["colunas"]] == ["registro", "nome"]


def test_parametros_invalidos(cliente):
    como("rh")
    assert cliente.get("/rh/relatorios/vinculos?inventado=1").status_code == 422
    assert cliente.get("/rh/relatorios/vinculos?colunas=nome,xyz").status_code == 422
    assert cliente.get("/rh/relatorios/vinculos?ordenar=xyz").status_code == 422
    assert cliente.get("/rh/relatorios/diagnostico-regra").status_code == 422          # regra obrigatória
    assert cliente.get("/rh/relatorios/diagnostico-regra?regra=X99").status_code == 422
    assert cliente.get("/rh/relatorios/afastamentos-retornos?dias=999").status_code == 422


def test_paginacao(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/vinculos?por_pagina=2&pagina=3").json()
    assert r["total"] == 5 and len(r["itens"]) == 1 and r["pagina"] == 3


# ============================================================================ diagnóstico
def test_diagnostico_regra_de_dependente(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/diagnostico-regra?regra=D03").json()
    assert r["total"] == 1
    assert r["contexto"]["regra"]["codigo"] == "D03"
    ids = [c["id"] for c in r["colunas"]]
    assert "dependente" in ids and ids[-1] == "detalhe"
    assert r["itens"][0]["nascimento"] == "2001-05-10"


def test_diagnostico_regra_filtra_por_secretaria_pelo_cadastro(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/diagnostico-regra?regra=S10&secretaria=000003").json()
    assert r["total"] == 2
    r = cliente.get("/rh/relatorios/diagnostico-regra?regra=S10&secretaria=000010").json()
    assert r["total"] == 0


def test_diagnostico_vinculos(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/diagnostico-vinculos").json()      # padrão: com erro ou alerta
    assert {i["registro"] for i in r["itens"]} == {"028310", "022818", "029982"}
    primeiro = r["itens"][0]
    assert primeiro["registro"] == "022818" and primeiro["erros"] == 2 and primeiro["regras"] == "D03, S10"
    r = cliente.get("/rh/relatorios/diagnostico-vinculos?situacao_cadastral=ok").json()
    assert {i["registro"] for i in r["itens"]} == {"028312", "016074"}
    r = cliente.get("/rh/relatorios/diagnostico-vinculos?regra=S01").json()
    assert [i["registro"] for i in r["itens"]] == ["029982"]


# ============================================================================ afastamentos
def test_afastamentos_vigentes(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/afastamentos-vigentes").json()
    assert r["total"] == 4 and r["data_referencia"] == "2026-09-29"
    item = next(i for i in r["itens"] if i["registro"] == "028310")
    assert item["secretaria"] == "SECRETARIA MUNICIPAL DE EDUCAÇÃO" and item["status_esocial"] == "Sem envio"
    assert {o["valor"] for o in r["opcoes"]["tipo"]} == {"12", "168", "378"}
    r = cliente.get("/rh/relatorios/afastamentos-vigentes?esocial=sem_status").json()
    assert r["total"] == 2
    r = cliente.get("/rh/relatorios/afastamentos-vigentes?sem_remuneracao=true").json()
    assert [i["registro"] for i in r["itens"]] == ["022818"]
    r = cliente.get("/rh/relatorios/afastamentos-vigentes", params={"q": "natalia"}).json()
    assert [i["registro"] for i in r["itens"]] == ["029982"]
    r = cliente.get("/rh/relatorios/afastamentos-vigentes?vinculo=02").json()     # completado pelo cadastro
    assert r["total"] == 4


def test_inconsistencias_colunas_por_regra(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/afastamentos-inconsistencias?regra=sobrepostos").json()
    assert r["itens"][0]["a_tipo"] == "12 LIC" and r["itens"][0]["dias_sobrepostos"] == 1
    assert [c["id"] for c in r["colunas"]] == fontes.COLUNAS_INCONSISTENCIA["sobrepostos"]


PARAMETROS_OBRIGATORIOS = {"diagnostico-regra": {"regra": "S10"}, "afastamentos-inconsistencias": {"regra": "sem-tipo"}}


@pytest.mark.parametrize("rid", list(POR_ID))
def test_todos_os_relatorios_respondem_e_exportam(cliente, rid):
    como("rh")
    params = PARAMETROS_OBRIGATORIOS.get(rid, {})
    r = cliente.get(f"/rh/relatorios/{rid}", params=params)
    assert r.status_code == 200, r.text
    corpo = r.json()
    assert corpo["relatorio"]["id"] == rid
    for formato in ("csv", "xlsx", "pdf"):
        e = cliente.get(f"/rh/relatorios/{rid}/exportar", params={**params, "formato": formato})
        assert e.status_code == 200, (formato, e.text)


# ============================================================================ exportação e auditoria
def test_exportar_csv_igual_a_tela(cliente):
    como("rh")
    e = cliente.get("/rh/relatorios/afastamentos-vigentes/exportar",
                    params={"q": "odenise", "colunas": "registro,nome,cpf,dtini,sem_remuneracao"})
    texto = e.content.decode("utf-8")
    assert texto.startswith("﻿")
    linhas = texto.lstrip("﻿").strip().splitlines()
    assert linhas[0] == "REGISTRO;Nome;CPF;Início;Sem remuneração"
    assert linhas[1:] == ["028310;ODENISE JARA GOMES LENTE;53175611168;01/06/2026;Não",
                          "028312;ODENISE JARA GOMES LENTE;53175611168;01/06/2026;Não"]
    assert "attachment" in e.headers["content-disposition"]


def test_exportar_xlsx(cliente):
    from openpyxl import load_workbook
    como("rh")
    e = cliente.get("/rh/relatorios/vinculos/exportar", params={"formato": "xlsx", "secretaria": EDU})
    ws = load_workbook(io.BytesIO(e.content)).active
    assert ws["A1"].value == "Vínculos do SIP"
    assert "Secretaria: SECRETARIA MUNICIPAL DE EDUCAÇÃO" in ws["A3"].value
    cab = [c.value for c in ws[6]]
    assert cab[0] == "REGISTRO" and "CPF" in cab
    assert ws.max_row == 6 + 3


def test_auditoria_registra_consulta_e_exportacao(cliente):
    como("rh")
    cliente.get("/rh/relatorios/vinculos", params={"q": "natalia", "secretaria": "000020"})
    cliente.get("/rh/relatorios/vinculos/exportar", params={"formato": "csv"})
    acoes = [(a["acao"], a["alvo"]) for a in AUDITORIA]
    assert acoes == [("relatorios.consultar", "vinculos"), ("relatorios.exportar", "vinculos")]
    assert AUDITORIA[0]["detalhe"]["q"] == "natalia" and AUDITORIA[0]["detalhe"]["total"] == 1
    assert AUDITORIA[0]["detalhe"]["filtros"]["secretaria"] == ["000020"]


def test_opcoes_contam_o_resultado_da_busca_e_dos_outros_filtros(cliente):
    como("rh")
    r = cliente.get("/rh/relatorios/afastamentos-vigentes", params={"q": "natalia"}).json()
    assert r["total"] == 1
    assert [(o["rotulo"], o["total"]) for o in r["opcoes"]["tipo"]] == [("LICENCA MATERNIDADE", 1)]
    # o próprio filtro não reduz as suas opções (dá para marcar outro valor)
    r = cliente.get("/rh/relatorios/afastamentos-vigentes?tipo=12").json()
    assert r["total"] == 2 and {o["valor"] for o in r["opcoes"]["tipo"]} == {"12", "168", "378"}
    assert {o["valor"]: o["total"] for o in r["opcoes"]["secretaria"]} == {"000003": 2}
    r = cliente.get("/rh/relatorios/afastamentos-vigentes?opcoes=false").json()
    assert r["opcoes"] == {} and r["total"] == 4


def test_exportar_pdf(cliente):
    como("rh")
    e = cliente.get("/rh/relatorios/afastamentos-vigentes/exportar", params={"formato": "pdf", "q": "odenise"})
    assert e.status_code == 200 and e.headers["content-type"] == "application/pdf"
    assert e.content.startswith(b"%PDF")
    from pypdf import PdfReader
    texto = PdfReader(io.BytesIO(e.content)).pages[0].extract_text()
    assert "Desenvolvimento — CTI Coordenadoria de Tecnologia da Informação" in texto
    assert "ODENISE" in texto and "LGPD" in texto
    assert AUDITORIA[-1]["detalhe"]["formato"] == "pdf"


def test_pdf_grande_demais(cliente, monkeypatch):
    from relatorios import exportar as EX
    monkeypatch.setattr(EX, "PDF_MAX", 2)
    como("rh")
    assert cliente.get("/rh/relatorios/afastamentos-vigentes/exportar?formato=pdf").status_code == 422

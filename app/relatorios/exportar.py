"""
Exportação do relatório filtrado: CSV (separador ";", UTF-8 com BOM, datas DD/MM/AAAA)
e Excel (.xlsx, com cabeçalho fixo, filtro automático e datas como data).
"""
import csv
import io
from datetime import date, datetime

def texto(v, tipo: str = "texto") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        return f"{v:.2f}".replace(".", ",")
    if isinstance(v, dict):
        return "; ".join(f"{k}: {x}" for k, x in v.items())
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    return str(v)


def para_csv(resultado: dict) -> bytes:
    colunas = resultado["colunas"]
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow([c["rotulo"] for c in colunas])
    for item in resultado["itens"]:
        w.writerow([texto(item.get(c["id"]), c["tipo"]) for c in colunas])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _celula(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, (int, float, date)):
        return v
    if isinstance(v, (dict, list)):
        return texto(v)
    return str(v)


def para_xlsx(resultado: dict, descricao_filtros: str) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as erro:  # pragma: no cover
        raise RuntimeError("Instale openpyxl no backend para exportar em Excel.") from erro

    colunas = resultado["colunas"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"[:31]
    ref = resultado.get("data_referencia")
    ws.append([resultado["relatorio"]["titulo"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Dados do SIP de {texto(ref)}" if ref else
               f"Gerado em {texto(resultado.get('gerado_em'))}"])
    ws.append([descricao_filtros or "Sem filtros"])
    ws.append([f"{resultado['total']} linha(s)"])
    ws.append([])
    linha_cab = ws.max_row + 1
    ws.append([c["rotulo"] for c in colunas])
    for celula in ws[linha_cab]:
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor="DDE6F0")
    for item in resultado["itens"]:
        ws.append([_celula(item.get(c["id"])) for c in colunas])

    for i, c in enumerate(colunas, start=1):
        letra = get_column_letter(i)
        if c["tipo"] == "data":
            for (cel,) in ws.iter_rows(min_row=linha_cab + 1, min_col=i, max_col=i):
                cel.number_format = "DD/MM/YYYY"
        if c["tipo"] == "percentual":
            for (cel,) in ws.iter_rows(min_row=linha_cab + 1, min_col=i, max_col=i):
                cel.number_format = "0.0"
        maior = max([len(str(c["rotulo"]))] + [len(texto(it.get(c["id"]))) for it in resultado["itens"][:500]])
        ws.column_dimensions[letra].width = min(max(10, maior + 2), 60)
    ws.freeze_panes = ws.cell(row=linha_cab + 1, column=1)
    ws.oddFooter.left.text = RODAPE_FIXO
    ws.oddFooter.right.text = "Página &P de &N"
    ws.page_setup.orientation = "landscape"
    if colunas:
        ws.auto_filter.ref = f"A{linha_cab}:{get_column_letter(len(colunas))}{max(ws.max_row, linha_cab)}"
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# --------------------------------------------------------------------------- PDF
PDF_MAX = 10_000   # acima disso o PDF fica impraticável: use Excel ou CSV
RODAPE_FIXO = "Desenvolvimento — CTI Coordenadoria de Tecnologia da Informação"


def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


CURTAS = {"registro", "matricula", "contrato", "cpf", "cpf_dependente", "reduzido", "item", "idade",
          "mes", "motivo_codigo", "faixa", "secretaria_codigo", "codigo"}
CODIGOS = {"registro", "matricula", "contrato", "reduzido", "item", "codigo", "secretaria_codigo", "motivo_codigo"}


def formatar(v, c: dict) -> str:
    """Texto de uma célula no PDF (pt-BR)."""
    if v is None or v == "":
        return "—"
    if c["tipo"] == "percentual" and isinstance(v, (int, float)):
        return f"{v:.1f}%".replace(".", ",")
    if isinstance(v, int) and not isinstance(v, bool) and c["id"] not in CODIGOS:
        return f"{v:,}".replace(",", ".")
    if isinstance(v, float):
        return f"{v:.2f}".replace(".", ",")
    if c["id"] in ("cpf", "cpf_dependente") and str(v).isdigit() and len(str(v)) == 11:
        d = str(v)
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    return texto(v)


def para_pdf(resultado: dict, descricao_filtros: str, usuario: str | None = None) -> bytes:
    """A4 deitado, cabeçalho repetido em cada página, rodapé com página, autor e aviso de LGPD."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    except ImportError as erro:  # pragma: no cover
        raise RuntimeError("Instale reportlab no backend para exportar em PDF.") from erro

    colunas = resultado["colunas"]
    itens = resultado["itens"]
    n = max(len(colunas), 1)
    corpo = 8 if n <= 8 else 7 if n <= 12 else 6

    s_titulo = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=13, leading=16)
    s_info = ParagraphStyle("i", fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    s_cab = ParagraphStyle("c", fontName="Helvetica-Bold", fontSize=corpo, leading=corpo + 2)
    s_txt = ParagraphStyle("x", fontName="Helvetica", fontSize=corpo, leading=corpo + 2)
    s_num = ParagraphStyle("n", parent=s_txt, alignment=TA_RIGHT)

    pagina = landscape(A4)
    margem = 10 * mm
    largura_util = pagina[0] - 2 * margem

    # Larguras: colunas curtas (datas, números, CPF, códigos) com a largura exata do conteúdo, sem quebrar;
    # as de texto dividem o resto, proporcional ao tamanho do texto.
    letra = corpo * 0.56            # largura média de um caractere em pontos
    folga = 8

    def maior_palavra(t: str) -> int:
        return max((len(p) for p in t.split()), default=0)

    fixas, pesos, minimos = {}, {}, {}
    for k, c in enumerate(colunas):
        valores = [formatar(i.get(c["id"]), c) for i in itens[:500]]
        cab = maior_palavra(c["rotulo"]) * 1.12   # cabeçalho em negrito é mais largo
        if c["tipo"] in ("data", "numero", "percentual", "sim_nao") or c["id"] in CURTAS:
            fixas[k] = max([cab] + [len(v) for v in valores]) * letra + folga
        else:
            maior = max([len(c["rotulo"])] + [len(v) for v in valores])
            pesos[k] = min(max(maior, 8), 50)
            minimos[k] = max([cab] + [maior_palavra(v) for v in valores]) * letra + folga
    total_fixas = sum(fixas.values())
    resto = largura_util - total_fixas
    soma = sum(pesos.values()) or 1
    larguras = []
    for k in range(len(colunas)):
        if k in fixas:
            larguras.append(fixas[k])
        else:
            larguras.append(max(minimos[k], resto * pesos[k] / soma))
    escala = largura_util / sum(larguras) if larguras and sum(larguras) > largura_util else 1
    larguras = [w * escala for w in larguras]

    def celula(v, c):
        t = _esc(formatar(v, c)).replace(" | ", "<br/>")
        return Paragraph(t, s_num if c["tipo"] in ("numero", "percentual") else s_txt)

    dados = [[Paragraph(_esc(c["rotulo"]), s_cab) for c in colunas]]
    dados += [[celula(i.get(c["id"]), c) for c in colunas] for i in itens]
    totais = resultado.get("totais")
    if totais:
        dados.append([Paragraph(f"<b>{_esc(texto(totais.get(c['id'])))}</b>", s_txt) for c in colunas])

    tabela = LongTable(dados, colWidths=larguras, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE6F0")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#8899AA")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
    ]
    for linha in range(2, len(dados), 2):
        estilo.append(("BACKGROUND", (0, linha), (-1, linha), colors.HexColor("#F5F7FA")))
    if totais:
        estilo.append(("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.black))
    tabela.setStyle(TableStyle(estilo))

    ref = resultado.get("data_referencia")
    gerado = datetime.now().strftime("%d/%m/%Y %H:%M")
    linhas_info = [
        f"Dados do SIP de {texto(ref)}" if ref else f"Diagnóstico gerado em {texto(resultado.get('gerado_em'))}",
        descricao_filtros or "Sem filtros",
        f"{resultado['total']} linha(s)" + (" · agrupado por pessoa" if resultado.get("agrupado_por_pessoa") else ""),
    ]
    historia = [Paragraph(_esc(resultado["relatorio"]["titulo"]), s_titulo), Spacer(1, 2 * mm)]
    historia += [Paragraph(_esc(t), s_info) for t in linhas_info]
    historia += [Spacer(1, 4 * mm), tabela]
    if not itens:
        historia.append(Paragraph("Nenhuma linha com esses filtros.", s_info))

    nominal = resultado["relatorio"].get("natureza") == "nominal"
    rodape_txt = f"Prefeitura de Cáceres · Pré-Censo · gerado em {gerado}" + (f" por {usuario}" if usuario else "")
    aviso = "Contém dados pessoais (LGPD): uso restrito ao serviço, não compartilhe." if nominal else ""

    def rodape(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        if aviso:
            canvas.setFont("Helvetica-Bold", 7)
            canvas.drawString(margem, 13 * mm, aviso)
        canvas.setStrokeColor(colors.HexColor("#BBBBBB"))
        canvas.setLineWidth(0.5)
        canvas.line(margem, 11 * mm, pagina[0] - margem, 11 * mm)
        canvas.setFont("Helvetica-Bold", 7)
        canvas.setFillColor(colors.HexColor("#333333"))
        canvas.drawString(margem, 7.5 * mm, RODAPE_FIXO)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(margem, 4 * mm, rodape_txt)
        canvas.drawRightString(pagina[0] - margem, 7.5 * mm, f"Página {doc.page}")
        canvas.restoreState()

    saida = io.BytesIO()
    doc = SimpleDocTemplate(saida, pagesize=pagina, leftMargin=margem, rightMargin=margem,
                            topMargin=margem, bottomMargin=18 * mm,
                            title=resultado["relatorio"]["titulo"], author=usuario or "")
    doc.build(historia, onFirstPage=rodape, onLaterPages=rodape)
    return saida.getvalue()
